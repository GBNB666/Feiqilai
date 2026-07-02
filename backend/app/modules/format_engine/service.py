import re
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from app.modules.format_engine.matcher import match_paragraph, is_in_reference_zone, is_reference_zone_end, is_equation_paragraph
from app.modules.format_engine.applier import (
    apply_page_settings,
    apply_format,
    apply_reference_format,
    apply_header_footer,
    apply_footer_page_number,
    apply_table_format,
    apply_equation_format,
    renumber_captions,
    renumber_references,
    should_skip_paragraph,
)
from app.modules.format_engine.toc import generate_toc


class FormatEngine:
    @staticmethod
    def execute(
        input_path: str,
        output_path: str,
        structure: dict,
        custom_settings: dict | None = None,
    ) -> None:
        doc = Document(input_path)
        matched_sections: set = set()

        # 提取自定义覆盖
        page_custom = custom_settings.get("page") if custom_settings else None
        flags = {
            "latin_font": custom_settings.get("latin_font", True) if custom_settings else True,
        }

        # 1. Page settings
        apply_page_settings(doc, page_custom)

        # 1.2 Pre-scan: set outlineLvl on headings BEFORE TOC generation,
        #     so generate_toc() can find them and build static entries.
        _prescan_set_outline_levels(doc, structure)

        # 1.5. TOC generation (if enabled) — now can detect headings via outlineLvl
        toc_skip_count = 0
        if custom_settings and custom_settings.get("toc_enabled"):
            toc_skip_count = generate_toc(doc, structure, custom_settings, flags)

        # 2. Paragraph matching and formatting loop
        in_references = False
        in_toc_zone = False
        title_applied = False
        para_abs_idx = 0  # absolute paragraph index in document body
        current_equation_chapter = 1
        equation_index = 0

        for para in doc.paragraphs:
            if should_skip_paragraph(para):
                continue

            # 纯图片段落（无文字）→ 仅设置间距，保持已有对齐（由图片格式化器控制）
            if _has_image(para):
                if not para.text.strip():
                    from docx.shared import Pt as Pt2
                    pf = para.paragraph_format
                    pf.space_before = Pt2(5)
                    pf.space_after = Pt2(5)
                continue

            text = para.text.strip()

            # ── Generated TOC paragraphs: skip from matching, apply TOC format directly ──
            if para_abs_idx < toc_skip_count:
                para_abs_idx += 1
                if text:
                    if para_abs_idx == 1:
                        apply_format(para, "toc_title", 0, custom_settings, flags)
                    else:
                        apply_format(para, "toc_entry", 0, custom_settings, flags)
                continue

            para_abs_idx += 1

            if not text:
                continue

            # TOC zone detection (existing TOC in original document)
            if _is_toc_heading(text) and not in_toc_zone and not title_applied:
                in_toc_zone = True
                apply_format(para, "toc_title", 0, custom_settings, flags)
                continue

            if in_toc_zone:
                if _is_body_start(text, structure):
                    in_toc_zone = False
                    # fall through to normal matching
                else:
                    apply_format(para, "toc_entry", 0, custom_settings, flags)
                    continue

            # ── Equation detection ──
            is_eq = is_equation_paragraph(text)
            if is_eq and not in_references:
                # 跟踪章节号（从最近的标题推断）
                eq_chapter = _get_current_equation_chapter(text, structure, matched_sections)
                if eq_chapter != current_equation_chapter:
                    current_equation_chapter = eq_chapter
                    equation_index = 0
                equation_index += 1
                eq_number = f"({current_equation_chapter}-{equation_index})"
                apply_equation_format(para, eq_number, custom_settings, flags)
                continue

            # Check reference zone
            in_references = is_in_reference_zone(text, in_references)

            if in_references:
                is_ref_heading = any(kw in text.replace(" ", "") for kw in ["参考文献", "References"])
                if is_ref_heading and "__references__" not in matched_sections:
                    # 参考文献标题本身 — 按标题格式
                    matched_sections.add("__references__")
                    apply_format(para, "heading", 1, custom_settings, flags)
                elif is_reference_zone_end(text, structure):
                    # 不像参考文献条目 → 已离开参考文献区域
                    in_references = False
                    # fall through to normal matching below
                else:
                    apply_reference_format(para, custom_settings)
                    continue

            # Match and apply
            match_type, level = match_paragraph(text, structure, matched_sections)

            if match_type == "paper_title" and title_applied:
                match_type = "body"
            elif match_type == "paper_title":
                title_applied = True

            apply_format(para, match_type, level, custom_settings, flags)

        # 3. Header/footer — style only (no content generation)
        apply_header_footer(doc, custom_settings, flags)
        apply_footer_page_number(doc, custom_settings)

        # 4. Table formatting: three-line borders + text style
        apply_table_format(doc, custom_settings, flags)

        # 5. Caption renumbering: 图X.Y / 表X.Y per chapter
        renumber_captions(doc, structure)

        # 6. Reference renumbering: continuous [1], [2], [3]...
        renumber_references(doc)

        # 7. Caption repositioning: 表注移到表格上方，图注移到图片下方
        _reposition_captions(doc)

        # 8. Save
        doc.save(output_path)

    @staticmethod
    def preview_formatted(
        input_path: str,
        structure: dict,
    ) -> list[dict]:
        """预览结构化的格式化结果，返回按章节聚合的段落列表。"""
        doc = Document(input_path)
        matched_sections: set = set()
        sections = structure.get("sections", [])

        # Init section buckets
        buckets: list[dict] = []
        bucket_map: dict[str, int] = {}
        for i, sec in enumerate(sections):
            numbering = sec.get("numbering", "")
            title = sec.get("title", "")
            key = f"{numbering}{title}"
            buckets.append({
                "level": sec.get("level", 0),
                "title": key,
                "content": [],
                "markers": [],
                "is_toc": False,
            })
            bucket_map[key] = i

        current_bucket = -1
        in_references = False

        for para in doc.paragraphs:
            if should_skip_paragraph(para):
                continue

            text = para.text.strip()
            if not text:
                continue

            in_references = is_in_reference_zone(text, in_references)

            prev_matched = set(matched_sections)
            match_type, level = match_paragraph(text, structure, matched_sections)

            if match_type == "heading":
                # 找到本次新增的 section key → 切换到对应 bucket
                new_keys = matched_sections - prev_matched
                for full_key in new_keys:
                    # 重建原始 key（含空格）来匹配 bucket_map
                    for sec in sections:
                        num = sec.get("numbering", "")
                        t = sec.get("title", "")
                        orig_key = f"{num}{t}"
                        clean_key = orig_key.strip().replace(" ", "")
                        if clean_key == full_key and orig_key in bucket_map:
                            current_bucket = bucket_map[orig_key]
                            break
                    if current_bucket >= 0:
                        break

            if current_bucket >= 0 and match_type not in ("paper_title", "heading"):
                buckets[current_bucket]["content"].append(text)

        return buckets


def _prescan_set_outline_levels(doc, structure: dict) -> None:
    """预扫文档，为所有匹配到的标题段落设置 outlineLvl。

    必须在 generate_toc() 之前调用，否则 TOC 找不到任何标题。
    使用独立的 matched_sections 集合，不影响后续主循环的匹配。
    """
    from app.modules.format_engine.matcher import match_paragraph
    from app.modules.format_engine.applier import _set_outline_level

    pre_matched: set = set()
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # 跳过表格内段落
        parent_tag = para._element.getparent().tag if para._element.getparent() is not None else ""
        parent_local = parent_tag.split("}")[-1] if "}" in parent_tag else parent_tag
        if parent_local == "tc":
            continue
        # 跳过图片段落
        if _has_image(para):
            continue

        match_type, level = match_paragraph(text, structure, pre_matched)

        if match_type == "heading" and level is not None:
            _set_outline_level(para, level - 1)  # 0-based


def _get_current_equation_chapter(text: str, structure: dict, matched_sections: set) -> int:
    """从已匹配的章节标题推断当前公式所属章节号。"""
    sections = structure.get("sections", [])
    clean = text.strip().replace(" ", "")
    for sec in reversed(sections):
        num = sec.get("numbering", "")
        t = sec.get("title", "")
        full = f"{num}{t}".strip().replace(" ", "")
        if full in matched_sections:
            ch_match = re.match(r"^(?:第)?(\d+)", num.strip())
            if ch_match:
                return int(ch_match.group(1))
    return 1


def _is_toc_heading(text: str) -> bool:
    """判断段落是否为目录标题（含"目录"且不超过 6 字）。"""
    cleaned = text.replace(" ", "").replace("　", "")
    return "目录" in cleaned and len(cleaned) <= 6


def _is_body_start(text: str, structure: dict) -> bool:
    """判断是否已离开目录区域进入正文。"""
    sections = structure.get("sections", [])
    if not sections:
        return False
    first_heading = ""
    for sec in sections:
        if sec.get("level") == 1:
            num = sec.get("numbering", "")
            title = sec.get("title", "")
            first_heading = f"{num}{title}".replace(" ", "")
            break
    if first_heading and first_heading in text.replace(" ", ""):
        return True
    keywords = ["摘要", "引言", "绪论", "Abstract", "Introduction"]
    for kw in keywords:
        if kw in text.replace(" ", ""):
            return True
    return False


# ── Caption repositioning ──────────────────────────────────────────

def _get_caption_type_simple(text: str) -> str | None:
    """Lightweight caption type check (same logic as matcher._get_caption_type)."""
    stripped = text.strip()
    if len(stripped) > 50:
        return None
    if re.match(r"^(?:表|Table\.?)\s*\d+", stripped):
        return "table_caption"
    if re.match(r"^(?:图|Fig(?:ure)?\.?)\s*\d+", stripped):
        return "figure_caption"
    return None


def _has_image(para) -> bool:
    """Check if a paragraph contains an inline image/drawing.
    Handles both namespace-qualified and unqualified element names."""
    # Search for drawings (namespace-qualified and plain)
    for child in para._element.iter():
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag in ("drawing", "pict"):
            return True
    return False


def _reposition_captions(doc) -> None:
    """遍历文档 body 子元素，确保：
    - table_caption 段落紧邻在 <w:tbl> 表格之前（表头在上）
    - figure_caption 段落紧邻在含图片的段落之后（图注在下）

    每次只移动一个元素，移动后重新扫描整个 body。
    """
    body = doc.element.body
    max_iterations = 20  # 安全上限

    for _ in range(max_iterations):
        children = list(body)
        moved = False

        # 构建当前 items 列表
        items = []
        for i, el in enumerate(children):
            tag = el.tag.split("}")[-1] if "}" in el.tag else el.tag
            if tag == "p":
                para = Paragraph(el, body)
                items.append((i, el, "p", para.text.strip(), para))
            elif tag == "tbl":
                items.append((i, el, "tbl", None, None))

        # ── 处理表注：如果在表格后面，移到表格前面 ──
        for idx, (orig_i, el, kind, text, para) in enumerate(items):
            if kind != "p" or not text:
                continue
            if _get_caption_type_simple(text) != "table_caption":
                continue

            # 检查前面紧邻的非空元素是不是 table
            prev_is_table = False
            for j in range(idx - 1, -1, -1):
                if items[j][2] == "tbl":
                    prev_is_table = True
                    table_el = items[j][1]
                    break
                if items[j][2] == "p" and items[j][3]:
                    break  # 中间有正文，不关联

            if not prev_is_table:
                continue  # 表注前面不是表格，位置正确（表注可能已在表格前）

            # 表注在表格后面 → 移到表格之前
            body.remove(el)
            table_pos = list(body).index(table_el)
            body.insert(table_pos, el)
            moved = True
            break  # 重新扫描

        if moved:
            continue

        # ── 处理图注：如果在图片前面，移到图片后面 ──
        for idx, (orig_i, el, kind, text, para) in enumerate(items):
            if kind != "p" or not text:
                continue
            if _get_caption_type_simple(text) != "figure_caption":
                continue

            # 检查前面紧邻的元素是不是图片段落
            prev_is_image = False
            for j in range(idx - 1, -1, -1):
                if items[j][2] == "p" and items[j][4] is not None:
                    if _has_image(items[j][4]):
                        prev_is_image = True
                        break
                if items[j][2] == "p" and items[j][3]:
                    if not (items[j][4] and _has_image(items[j][4])):
                        break  # 中间有正文，不关联

            if prev_is_image:
                continue  # 图片已在图注前面，位置正确

            # 找后面最近的图片段落
            next_img_idx = None
            next_img_el = None
            for j in range(idx + 1, len(items)):
                if items[j][2] == "p" and items[j][4] is not None:
                    if _has_image(items[j][4]):
                        next_img_idx = items[j][0]
                        next_img_el = items[j][1]
                        break
                if items[j][2] == "p" and items[j][3]:
                    break  # 中间有正文

            if next_img_el is None:
                continue

            # 图注在图片前面 → 移到图片后面
            body.remove(el)
            img_pos = list(body).index(next_img_el)
            body.insert(img_pos + 1, el)
            moved = True
            break

        if not moved:
            break  # 无需更多移动

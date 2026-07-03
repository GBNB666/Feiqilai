# -*- coding: utf-8 -*-
"""Generate business plan docx - reads content from markdown draft."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os, re, json

BASE = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(BASE, "docs", "ZhiPaiAI_Business_Plan.docx")

doc = Document()

for sec in doc.sections:
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

def sf(run, cn, en, sz, bold=False):
    run.font.size = Pt(sz)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:eastAsia'), cn)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:cs'), en)

def ap(text, cn="KaiTi", en="Times New Roman", sz=14, bold=False, al="distribute", sb=0, sa=0, ls=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    sf(r, cn, en, sz, bold)
    am = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
          "distribute": WD_ALIGN_PARAGRAPH.DISTRIBUTE, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    p.paragraph_format.alignment = am.get(al, WD_ALIGN_PARAGRAPH.DISTRIBUTE)
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    if ls:
        p.paragraph_format.line_spacing = Pt(ls)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    return p

def h1(t): ap(t, "HeiTi", "Times New Roman", 18, False, "left", 9, 9)
def h2(t): ap(t, "KaiTi", "Times New Roman", 16, True, "left", 8, 8)
def body(t): ap(t, "KaiTi", "Times New Roman", 14, False, "distribute", 0, 0, 22)

print("Generating business plan...")
print("Font names: KaiTi for body, HeiTi for headings")
print("Make sure KaiTi and HeiTi fonts are installed on this system")

# ── COVER ──
for _ in range(6): doc.add_paragraph()
ap("《AI智排》创业计划书", "KaiTi", "Times New Roman", 26, True, "center")
doc.add_paragraph()
ap("——基于AI的论文格式智能排版平台", "KaiTi", "Times New Roman", 16, False, "center")
for _ in range(4): doc.add_paragraph()
ap("项目发起人：博博", "KaiTi", "Times New Roman", 16, False, "center")
ap("所在院校：护理学院", "KaiTi", "Times New Roman", 16, False, "center")
ap("联系方式：_________", "KaiTi", "Times New Roman", 16, False, "center")
doc.add_page_break()

# ── TOC ──
ap("目  录", "KaiTi", "Times New Roman", 14, True, "center", 0, 20)
toc = []
for i in range(1, 9):
    cn_num = ["", "一", "二", "三", "四", "五", "六", "七", "八"][i]
    toc.append(f"{cn_num}、{['','概述','产品/服务','市场','竞争','营销','经营','组织','财务'][i]}")

for item in toc:
    ap(item, "KaiTi", "Times New Roman", 14, True, "left", 0, 0, 22)
doc.add_page_break()

# ── BODY from markdown ──
md_path = os.path.join(BASE, "docs", "创业计划书-草稿.md")
if not os.path.exists(md_path):
    print(f"ERROR: markdown draft not found at {md_path}")
    exit(1)

with open(md_path, 'r', encoding='utf-8') as f:
    md = f.read()

print(f"Read {len(md)} chars from draft")

# Parse: split by ## headings
blocks = re.split(r'\n(?=## )', md)
pending_table = None

for block in blocks:
    lines = block.strip().split('\n')
    if not lines: continue

    for i, line in enumerate(lines):
        s = line.strip()
        if not s: continue

        # Top-level title - skip
        if s.startswith('# ') and not s.startswith('## '):
            continue

        # H1: "## 一、概述"
        if s.startswith('## '):
            title = s[3:].strip()
            h1(title)
            continue

        # H2: "### 1.1 xxx"
        if s.startswith('### '):
            title = s[4:].strip()
            h2(title)
            continue

        # Table
        if s.startswith('|') and s.endswith('|'):
            cells = [c.strip() for c in s.split('|')[1:-1]]
            if all(re.match(r'^[-:\s]+$', c) for c in cells):
                continue  # separator row
            # Emit as text since table parsing is complex
            body(" | ".join(cells))
            continue

        # Horizontal rule
        if s in ('---', '***', '___'):
            continue

        # Clean markdown formatting
        clean = s
        clean = re.sub(r'\*\*(.+?)\*\*', r'\1', clean)
        clean = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', clean)
        clean = re.sub(r'`(.+?)`', r'\1', clean)
        # Remove leading list markers
        clean = re.sub(r'^[\d]+\.\s*', '', clean)

        body(clean)

# ── FOOTER ──
doc.add_paragraph()
ap("本计划书由博博撰写，基于真实项目paper-formatter论文AI排版系统的开发实践。项目已历经V1、V2两次迭代，V3正在开发中，始终坚持先跑通再美化的务实开发理念。", "KaiTi", "Times New Roman", 10, False, "center")

doc.save(OUTPUT)
print(f"Done: {OUTPUT}")

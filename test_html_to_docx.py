"""模拟"浏览器复制HTML → 粘贴到Word"的格式化保留效果"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from bs4 import BeautifulSoup
import re

doc = Document()

# 设置页面边距
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

with open("test-paper-formatted.html", "r", encoding="utf-8") as f:
    html = f.read()

soup = BeautifulSoup(html, 'html.parser')

def pt_to_emus(pt_val):
    """Convert points to EMU (for spacing)"""
    return Pt(pt_val)

def apply_style(para, run=None, font_name="宋体", font_size=12, bold=False,
                align="justify", line_spacing=1.5, first_indent=True):
    """Apply formatting to a paragraph and its runs."""
    if run is None:
        if para.runs:
            run = para.runs[0]
        else:
            run = para.add_run("")

    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold

    # Set East Asian font
    try:
        from lxml import etree
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if rFonts is None:
            rFonts = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font_name)
    except Exception:
        pass

    align_map = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    para.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    para.paragraph_format.line_spacing = line_spacing

    if first_indent:
        para.paragraph_format.first_line_indent = Pt(24)  # 2 chars at 12pt
    else:
        para.paragraph_format.first_line_indent = None

# Process elements in body
body = soup.find('body')
for el in body.find_all(recursive=False):
    tag = el.name
    cls = el.get('class', [])
    text = el.get_text(strip=False).strip()

    if not text:
        continue

    para = doc.add_paragraph()
    run = para.add_run(text)

    if 'paper-title' in cls or tag == 'h1' or 'ref-title' in cls:
        # 一级标题: 黑体 16pt 加粗 居中
        apply_style(para, run, font_name="黑体", font_size=16, bold=True,
                    align="center", line_spacing=1.5, first_indent=False)
    elif tag == 'h2':
        # 二级标题: 黑体 14pt 加粗 左对齐
        apply_style(para, run, font_name="黑体", font_size=14, bold=True,
                    align="left", line_spacing=1.5, first_indent=False)
    elif tag == 'h3':
        # 三级标题: 黑体 12pt 加粗 左对齐
        apply_style(para, run, font_name="黑体", font_size=12, bold=True,
                    align="left", line_spacing=1.5, first_indent=False)
    elif 'caption' in cls:
        # 图注/表注: 宋体 10.5pt 居中
        apply_style(para, run, font_name="宋体", font_size=10.5, bold=False,
                    align="center", line_spacing=1.5, first_indent=False)
    elif 'ref-item' in cls:
        # 参考文献条目: 宋体 10.5pt 左对齐
        apply_style(para, run, font_name="宋体", font_size=10.5, bold=False,
                    align="left", line_spacing=1.2, first_indent=False)
    elif 'thanks' in cls:
        apply_style(para, run, font_name="宋体", font_size=12, bold=False,
                    align="left", line_spacing=1.5, first_indent=False)
    else:
        # 正文: 宋体 12pt 两端对齐 1.5行距 首行缩进
        apply_style(para, run, font_name="宋体", font_size=12, bold=False,
                    align="justify", line_spacing=1.5, first_indent=True)

output_path = "test-paper-from-html.docx"
doc.save(output_path)
print(f"Done: {output_path}")
print(f"Open it to check if formatting is preserved")
